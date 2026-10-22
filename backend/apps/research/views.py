import io
import random
from pathlib import Path
from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Count, Avg, Q
from django.http import FileResponse
from django.shortcuts import get_object_or_404
from rest_framework import serializers, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError, PermissionDenied
from rest_framework.response import Response
from drf_spectacular.utils import extend_schema
from . import models as m, serializers as s, services, tasks
from .validation import inspect_image, check_boxes


def owned(queryset, user, field="owner"):
    if not user.is_authenticated:
        return queryset.none()
    return queryset if user.administrator else queryset.filter(**{field: user})


def file_response(field, filename=None, download=False):
    if not field:
        raise ValidationError("Archivo todavía no disponible.")
    return FileResponse(
        field.open("rb"),
        as_attachment=download,
        filename=filename or Path(field.name).name,
    )


def validated(serializer, data):
    obj = serializer(data=data)
    obj.is_valid(raise_exception=True)
    return obj.validated_data


class ReadOnly(viewsets.ReadOnlyModelViewSet):
    pass


class DatasetViewSet(viewsets.ModelViewSet):
    serializer_class = s.DatasetSerializer
    http_method_names = ["get", "post", "patch", "head", "options"]

    def get_queryset(self):
        qs = owned(m.Dataset.objects.all(), self.request.user)
        if self.request.query_params.get("search"):
            qs = qs.filter(name__icontains=self.request.query_params["search"])
        return qs

    def perform_create(self, serializer):
        obj = serializer.save(owner=self.request.user)
        services.audit(self.request.user, "dataset.create", obj)

    @action(detail=True, methods=["post"], url_path="import")
    def import_files(self, request, pk=None):
        files = request.FILES.getlist("files")
        if not files:
            raise ValidationError("Selecciona imágenes o un ZIP.")
        if sum(x.size for x in files) > 100 * 1024 * 1024:
            raise ValidationError("Carga máxima 100 MB.")
        return Response(
            services.import_images(
                self.get_object(),
                files,
                request.user,
                str(request.data.get("source", ""))[:500],
                str(request.data.get("fruit_code", ""))[:120],
            ),
            status=201,
        )

    @action(detail=True, methods=["post"])
    def split(self, request, pk=None):
        class Input(serializers.Serializer):
            ratios = serializers.ListField(
                child=serializers.FloatField(min_value=0.01, max_value=0.98),
                min_length=3,
                max_length=3,
                default=[0.7, 0.2, 0.1],
            )
            seed = serializers.IntegerField(
                min_value=0, max_value=2**31 - 1, default=42
            )

        data = validated(Input, request.data)
        if abs(sum(data["ratios"]) - 1) > 0.0001:
            raise ValidationError("Las proporciones deben sumar 1.")
        return Response(
            services.split_dataset(
                self.get_object(), data["ratios"], data["seed"], request.user
            )
        )

    @action(detail=True, methods=["post"])
    def freeze(self, request, pk=None):
        class Input(serializers.Serializer):
            seed = serializers.IntegerField(
                min_value=0, max_value=2**31 - 1, default=42
            )
            variants = serializers.DictField(
                child=serializers.UUIDField(), default=dict
            )
            reviewed_variants = serializers.BooleanField(default=False)

        data = validated(Input, request.data)
        if data["variants"] and not data["reviewed_variants"]:
            raise ValidationError(
                "Confirma la revisión de las cajas transformadas antes de usar derivados."
            )
        dataset = self.get_object()
        for image_id, variant_id in data["variants"].items():
            if not m.ImageVariant.objects.filter(
                pk=variant_id, image_id=image_id, image__dataset=dataset
            ).exists():
                raise ValidationError("Derivado inválido.")
        return Response(
            s.VersionSerializer(
                services.freeze(dataset, request.user, data["seed"], data["variants"])
            ).data,
            status=201,
        )

    @action(detail=True, methods=["get"])
    def duplicates(self, request, pk=None):
        images = list(
            self.get_object().images.only("id", "name", "perceptual_hash")[:2000]
        )
        pairs = []
        for i, a in enumerate(images):
            for b in images[i + 1 :]:
                distance = (
                    int(a.perceptual_hash, 16) ^ int(b.perceptual_hash, 16)
                ).bit_count()
                if distance <= 4:
                    pairs.append(
                        {
                            "first": str(a.pk),
                            "second": str(b.pk),
                            "first_name": a.name,
                            "second_name": b.name,
                            "distance": distance,
                        }
                    )
                    if len(pairs) >= 200:
                        return Response({"pairs": pairs, "truncated": True})
        return Response(
            {
                "pairs": pairs,
                "truncated": False,
                "note": "Sugerencias por dHash; revisar manualmente, no eliminaciones automáticas.",
            }
        )


class ImageViewSet(viewsets.GenericViewSet):
    serializer_class = s.ImageSerializer

    def get_queryset(self):
        qs = owned(
            m.ImageAsset.objects.select_related("dataset"),
            self.request.user,
            "dataset__owner",
        )
        for field in ["dataset", "state", "split"]:
            if self.request.query_params.get(field):
                qs = qs.filter(**{field: self.request.query_params[field]})
        if self.request.query_params.get("search"):
            qs = qs.filter(name__icontains=self.request.query_params["search"])
        return qs

    def list(self, request):
        page = self.paginate_queryset(self.get_queryset())
        return self.get_paginated_response(self.get_serializer(page, many=True).data)

    def retrieve(self, request, pk=None):
        return Response(self.get_serializer(self.get_object()).data)

    def partial_update(self, request, pk=None):
        obj = self.get_object()
        serializer = self.get_serializer(obj, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        services.audit(request.user, "image.metadata", obj)
        return Response(serializer.data)

    @action(detail=True, methods=["get"])
    def file(self, request, pk=None):
        return file_response(self.get_object().file)

    @action(detail=True, methods=["post"])
    def annotate(self, request, pk=None):
        data = validated(s.AnnotationInput, request.data)
        obj = services.annotate(
            self.get_object(),
            request.user,
            boxes=data["annotations"],
            state=data["state"],
            revision=data["revision"],
            note=data.get("note", ""),
        )
        return Response(self.get_serializer(obj).data)

    @action(detail=True, methods=["get"])
    def history(self, request, pk=None):
        return Response(
            list(
                self.get_object().history.values(
                    "revision",
                    "annotations",
                    "state",
                    "note",
                    "author__username",
                    "created_at",
                )
            )
        )


class VersionViewSet(ReadOnly):
    serializer_class = s.VersionSerializer

    def get_queryset(self):
        qs = owned(
            m.DatasetVersion.objects.select_related("dataset"),
            self.request.user,
            "dataset__owner",
        )
        if self.request.query_params.get("dataset"):
            qs = qs.filter(dataset_id=self.request.query_params["dataset"])
        return qs

    @action(detail=True, methods=["get"])
    def download(self, request, pk=None):
        version = self.get_object()
        services.verify_version_archive(version)
        return file_response(version.archive, "dataset-yolo.zip", True)


class VariantViewSet(ReadOnly):
    serializer_class = s.VariantSerializer

    def get_queryset(self):
        qs = owned(
            m.ImageVariant.objects.select_related("image__dataset"),
            self.request.user,
            "image__dataset__owner",
        )
        if self.request.query_params.get("image"):
            qs = qs.filter(image_id=self.request.query_params["image"])
        return qs

    @action(detail=True, methods=["get"])
    def file(self, request, pk=None):
        return file_response(self.get_object().file)


class JobViewSet(ReadOnly):
    serializer_class = s.JobSerializer

    def get_queryset(self):
        return owned(m.Job.objects.all(), self.request.user)

    @action(detail=False, methods=["post"])
    def preprocessing(self, request):
        data = validated(s.PreprocessingInput, request.data)
        images = owned(
            m.ImageAsset.objects.all(), request.user, "dataset__owner"
        ).filter(pk__in=data["images"])
        if images.count() != len(set(data["images"])):
            raise PermissionDenied("Imágenes no disponibles.")
        data["images"] = [str(x) for x in data["images"]]
        job = m.Job.objects.create(
            owner=request.user, kind="preprocessing", config=data
        )
        return Response(
            self.get_serializer(services.dispatch(job, tasks.preprocess)).data,
            status=202,
        )

    @action(detail=True, methods=["post"])
    def retry(self, request, pk=None):
        job = self.get_object()
        with transaction.atomic():
            job = m.Job.objects.select_for_update().get(pk=job.pk)
            if job.state not in ["queued", "failed"]:
                raise ValidationError(
                    "Solo se pueden reenviar trabajos en cola o fallidos."
                )
            # A task that completed its artifact but crashed before final status must not duplicate it.
            if (
                job.kind == "training"
                and m.AIModel.objects.filter(run__job=job).exists()
            ):
                raise ValidationError(
                    "Ya existe un modelo para este trabajo; revisar estado con administrador."
                )
            job.state = "queued"
            job.error = ""
            job.save()
        task = {
            "training": tasks.train,
            "inference": tasks.infer,
            "preprocessing": tasks.preprocess,
            "evaluation": tasks.validate_model,
        }[job.kind]
        return Response(
            self.get_serializer(services.dispatch(job, task)).data, status=202
        )


class TrainingViewSet(ReadOnly):
    serializer_class = s.TrainingSerializer

    def get_queryset(self):
        return owned(
            m.TrainingRun.objects.select_related("job"), self.request.user, "job__owner"
        )

    @extend_schema(request=s.TrainingInput, responses=s.TrainingSerializer)
    def create(self, request):
        data = validated(s.TrainingInput, request.data)
        version = get_object_or_404(
            owned(m.DatasetVersion.objects.all(), request.user, "dataset__owner"),
            pk=data["version"],
        )
        data["device"] = settings.AI_DEVICE if data["device"] == "server" else "cpu"
        data["version"] = str(data["version"])
        with transaction.atomic():
            job = m.Job.objects.create(owner=request.user, kind="training", config=data)
            run = m.TrainingRun.objects.create(
                version=version, job=job, architecture=data["architecture"]
            )
        services.dispatch(job, tasks.train)
        run.refresh_from_db()
        services.audit(request.user, "training.create", run)
        return Response(self.get_serializer(run).data, status=202)

    @action(detail=True, methods=["get"])
    def artifacts(self, request, pk=None):
        run = self.get_object()
        attempt = serializers.IntegerField(min_value=1).run_validation(
            request.query_params.get("attempt", max(1, run.attempt))
        )
        if attempt > max(1, run.attempt):
            raise ValidationError("Intento no disponible.")
        root = Path(settings.MEDIA_ROOT) / "training-artifacts" / str(run.pk)
        folder = root / f"attempt-{attempt}" / "fit"
        # Legacy experiments stored their artifacts directly under the run folder.
        if run.attempt == 0:
            folder = root
        allowed = {".png", ".csv", ".yaml", ".json", ".jsonl", ".pt", ".log"}
        filename = request.query_params.get("name")
        if filename:
            if Path(filename).name != filename or Path(filename).suffix not in allowed:
                raise ValidationError("Archivo inválido.")
            path = folder / filename
            if not path.is_file():
                raise ValidationError("Archivo no disponible.")
            return FileResponse(path.open("rb"), as_attachment=True, filename=filename)
        return Response(
            sorted(
                p.name for p in folder.iterdir() if p.is_file() and p.suffix in allowed
            )
            if folder.exists()
            else []
        )


class ModelViewSet(ReadOnly):
    serializer_class = s.ModelSerializer

    def get_queryset(self):
        qs = m.AIModel.objects.select_related("run")
        if not self.request.user.is_authenticated:
            return qs.none()
        return (
            qs
            if self.request.user.administrator
            else qs.filter(Q(run__job__owner=self.request.user) | Q(active=True))
        )

    @action(detail=True, methods=["post"])
    def activate(self, request, pk=None):
        if not request.user.administrator:
            raise PermissionDenied("Solo el administrador puede activar modelos.")
        obj = self.get_object()
        with transaction.atomic():
            # Lock a stable row even when there is currently no active model.
            from apps.accounts.models import User

            User.objects.select_for_update().get(pk=request.user.pk)
            list(m.AIModel.objects.select_for_update().order_by("id"))
            m.AIModel.objects.filter(active=True).update(active=False)
            obj.active = True
            obj.approved_by = request.user
            obj.save()
            services.audit(request.user, "model.activate", obj)
        return Response(self.get_serializer(obj).data)

    @action(detail=True, methods=["get"])
    def download(self, request, pk=None):
        return file_response(self.get_object().weights, "modelo-rona.pt", True)


class InspectionViewSet(ReadOnly):
    serializer_class = s.InspectionSerializer

    def get_queryset(self):
        qs = owned(
            m.Inspection.objects.select_related("job", "model"), self.request.user
        )
        if self.request.query_params.get("search"):
            qs = qs.filter(name__icontains=self.request.query_params["search"])
        for param, field in [
            ("from", "created_at__date__gte"),
            ("to", "created_at__date__lte"),
        ]:
            if self.request.query_params.get(param):
                value = serializers.DateField().run_validation(
                    self.request.query_params[param]
                )
                qs = qs.filter(**{field: value})
        return qs

    def create(self, request):
        class Input(serializers.Serializer):
            model = serializers.UUIDField()
            confidence = serializers.FloatField(
                min_value=0.01, max_value=0.99, default=0.25
            )

        data = validated(Input, request.data)
        models = m.AIModel.objects.all()
        if not request.user.administrator:
            models = models.filter(Q(run__job__owner=request.user) | Q(active=True))
        model = get_object_or_404(models, pk=data["model"])
        files = request.FILES.getlist("files")
        if (
            not files
            or len(files) > 100
            or sum(f.size for f in files) > 100 * 1024 * 1024
        ):
            raise ValidationError("Selecciona hasta 100 imágenes, máximo total 100 MB.")
        prepared = []
        for upload in files:
            contents = upload.read()
            meta = inspect_image(contents, upload.name)
            prepared.append((upload.name, contents, meta))
        inspections = []
        for name, contents, meta in prepared:
            with transaction.atomic():
                job = m.Job.objects.create(
                    owner=request.user,
                    kind="inference",
                    config={"confidence": data["confidence"]},
                )
                obj = m.Inspection(
                    owner=request.user,
                    model=model,
                    job=job,
                    name=Path(name).name,
                    sha256=meta["sha256"],
                )
                obj.original.save(
                    f"{obj.pk}{Path(name).suffix.lower()}",
                    ContentFile(contents),
                    save=False,
                )
                obj.save()
            services.dispatch(job, tasks.infer)
            obj.refresh_from_db()
            inspections.append(obj)
        return Response(self.get_serializer(inspections, many=True).data, status=202)

    @action(detail=True, methods=["get"])
    def file(self, request, pk=None):
        obj = self.get_object()
        return file_response(
            obj.annotated
            if request.query_params.get("annotated") == "1"
            else obj.original,
            download=request.query_params.get("download") == "1",
        )

    @action(detail=True, methods=["post"])
    def review(self, request, pk=None):
        class Input(serializers.Serializer):
            decision = serializers.ChoiceField(choices=["confirmed", "corrected"])
            note = serializers.CharField(max_length=4000)
            annotations = serializers.JSONField(default=list)

        data = validated(Input, request.data)
        data["annotations"] = check_boxes(data["annotations"])
        obj = self.get_object()
        if obj.job.state != "succeeded":
            raise ValidationError("La inspección debe estar completada.")
        with transaction.atomic():
            m.InspectionReview.objects.create(
                inspection=obj, author=request.user, **data
            )
            obj.review = data["decision"]
            obj.save(update_fields=["review"])
        return Response(self.get_serializer(obj).data)

    @action(detail=True, methods=["get"])
    def reviews(self, request, pk=None):
        return Response(
            list(
                self.get_object().reviews.values(
                    "decision", "note", "annotations", "author__username", "created_at"
                )
            )
        )


class EvaluationViewSet(ReadOnly):
    serializer_class = s.EvaluationSerializer

    def get_queryset(self):
        return owned(m.Evaluation.objects.all(), self.request.user)

    def create(self, request):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        if (
            not request.user.administrator
            and data["version"].dataset.owner_id != request.user.pk
        ):
            raise PermissionDenied()
        if (
            not request.user.administrator
            and data["model"].run.job.owner_id != request.user.pk
            and not data["model"].active
        ):
            raise PermissionDenied()
        test = [x for x in data["version"].manifest if x["split"] == "test"]
        training = {
            x["original_sha256"]
            for x in data["model"].run.version.manifest
            if x["split"] in ["train", "val"]
        }
        if not test or any(x["original_sha256"] in training for x in test):
            raise ValidationError(
                "Test vacío o contiene imágenes usadas para ajustar el modelo."
            )
        with transaction.atomic():
            obj = serializer.save(owner=request.user, locked=True)
            m.EvaluationSample.objects.bulk_create(
                [
                    m.EvaluationSample(
                        evaluation=obj,
                        image_id_snapshot=x["image_id"],
                        fruit_code=x["fruit_code"],
                        truth=bool(x["annotations"]),
                    )
                    for x in test
                ]
            )
        return Response(self.get_serializer(obj).data, status=201)

    @action(detail=True, methods=["get"])
    def samples(self, request, pk=None):
        evaluation = self.get_object()
        samples = list(evaluation.samples.all())
        random.Random(str(evaluation.pk) + str(request.user.pk)).shuffle(samples)
        # Blinded API omits reference and AI results.
        return Response(
            [
                {
                    "id": x.pk,
                    "fruit_code": x.fruit_code,
                    "image_id": x.image_id_snapshot,
                    "completed": list(
                        x.observations.values_list("participant", flat=True)
                    ),
                }
                for x in samples
            ]
        )

    @action(detail=True, methods=["get"], url_path="sample-file")
    def sample_file(self, request, pk=None):
        import zipfile

        evaluation = self.get_object()
        sample = get_object_or_404(
            evaluation.samples, pk=request.query_params.get("sample")
        )
        item = next(
            x
            for x in evaluation.version.manifest
            if x["image_id"] == sample.image_id_snapshot
        )
        with zipfile.ZipFile(evaluation.version.archive.path) as archive:
            data = archive.read(f"images/test/{item['file']}")
        return FileResponse(io.BytesIO(data), filename=item["file"])

    @action(detail=True, methods=["post"])
    def observe(self, request, pk=None):
        class Input(serializers.Serializer):
            sample = serializers.UUIDField()
            participant = serializers.CharField(max_length=80)
            positive = serializers.BooleanField()
            elapsed_ms = serializers.FloatField(min_value=1, max_value=86_400_000)

        data = validated(Input, request.data)
        evaluation = self.get_object()
        sample = get_object_or_404(evaluation.samples, pk=data.pop("sample"))
        if sample.observations.filter(participant=data["participant"]).exists():
            raise ValidationError(
                "Ya existe una observación de este participante para la muestra."
            )
        observation = m.ManualObservation.objects.create(
            sample=sample, observer=request.user, **data
        )
        services.audit(request.user, "evaluation.observe", observation)
        return Response({"id": observation.pk}, status=201)

    @action(detail=True, methods=["post"])
    def run(self, request, pk=None):
        evaluation = self.get_object()
        with transaction.atomic():
            evaluation = m.Evaluation.objects.select_for_update().get(pk=evaluation.pk)
            if evaluation.job_id:
                raise ValidationError(
                    "Esta evaluación ya tiene un trabajo; consulta su estado o reinténtalo desde Trabajos."
                )
            job = m.Job.objects.create(
                owner=request.user,
                kind="evaluation",
                config={"evaluation": str(evaluation.pk)},
            )
            evaluation.job = job
            evaluation.save(update_fields=["job"])
        return Response(
            s.JobSerializer(services.dispatch(job, tasks.validate_model)).data,
            status=202,
        )

    @action(detail=True, methods=["get"])
    def export(self, request, pk=None):
        from openpyxl import Workbook

        evaluation = self.get_object()
        wb = Workbook()
        ws = wb.active
        ws.title = "Muestras"
        ws.append(
            [
                "evaluacion",
                "muestra",
                "fruto",
                "referencia_rona",
                "sistema_rona",
                "tiempo_sistema_ms",
                "participante",
                "manual_rona",
                "tiempo_manual_ms",
            ]
        )

        def safe(value):
            return (
                "'" + value
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
                else value
            )

        for sample in evaluation.samples.prefetch_related("observations"):
            observations = list(sample.observations.all()) or [None]
            for obs in observations:
                ws.append(
                    [
                        safe(v)
                        for v in [
                            str(evaluation.pk),
                            sample.image_id_snapshot,
                            sample.fruit_code,
                            sample.truth,
                            sample.system_positive,
                            sample.system_ms,
                            obs.participant if obs else None,
                            obs.positive if obs else None,
                            obs.elapsed_ms if obs else None,
                        ]
                    ]
                )
        meta = wb.create_sheet("Protocolo")
        for key, value in {
            "nombre": evaluation.name,
            "protocolo": evaluation.protocol,
            "version_id": str(evaluation.version_id),
            "manifest_hash": evaluation.version.manifest_hash,
            "modelo_id": str(evaluation.model_id),
            "pesos_sha256": evaluation.model.sha256,
            "confianza": evaluation.confidence,
            "unidad": "imagen; resumen comparativo agregado por fruto",
            "tiempos": "no incluyen captura ni subida; sistema incluye renderizado de resultado",
            "fecha": evaluation.created_at.isoformat(),
        }.items():
            meta.append([key, safe(value)])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        services.audit(request.user, "evaluation.export", evaluation)
        return FileResponse(output, as_attachment=True, filename="evaluacion.xlsx")

    @action(detail=True, methods=["get"])
    def summary(self, request, pk=None):
        from ai.evaluation import paired

        evaluation = self.get_object()
        results = {}
        participants = (
            m.ManualObservation.objects.filter(sample__evaluation=evaluation)
            .values_list("participant", flat=True)
            .distinct()
        )
        for participant in participants:
            groups = {}
            samples = list(evaluation.samples.prefetch_related("observations"))
            for sample in samples:
                groups.setdefault(sample.fruit_code, []).append(sample)
            truth = []
            system = []
            manual = []
            st = []
            mt = []
            for fruit, samples in groups.items():
                observations = [
                    next(
                        (
                            o
                            for o in sample.observations.all()
                            if o.participant == participant
                        ),
                        None,
                    )
                    for sample in samples
                ]
                if any(o is None for o in observations) or any(
                    x.system_positive is None for x in samples
                ):
                    continue
                truth.append(any(x.truth for x in samples))
                system.append(any(x.system_positive for x in samples))
                manual.append(any(o.positive for o in observations))
                st.append(sum(x.system_ms for x in samples))
                mt.append(sum(o.elapsed_ms for o in observations))
            if truth:
                results[participant] = paired(truth, system, manual, st, mt)
        return Response(
            {
                "detection": evaluation.metrics,
                "participants": results,
                "note": "Se agregan vistas del mismo fruto con OR y se suman tiempos. Solo se comparan frutos con todas sus observaciones. No incluye captura/subida; el protocolo debe especificarlo.",
            }
        )


class DashboardViewSet(viewsets.ViewSet):
    def list(self, request):
        datasets = owned(m.Dataset.objects.all(), request.user)
        images = owned(m.ImageAsset.objects.all(), request.user, "dataset__owner")
        inspections = owned(m.Inspection.objects.all(), request.user)
        runs = owned(m.TrainingRun.objects.all(), request.user, "job__owner")
        active = m.AIModel.objects.filter(active=True).first()
        return Response(
            {
                "datasets": datasets.count(),
                "images": images.count(),
                "reviewed": images.filter(reviewed_by__isnull=False).count(),
                "inspections": inspections.count(),
                "training_runs": runs.count(),
                "mean_inference_ms": inspections.filter(
                    job__state="succeeded"
                ).aggregate(v=Avg("inference_ms"))["v"],
                "states": list(images.values("state").annotate(count=Count("id"))),
                "active_model": s.ModelSerializer(active).data if active else None,
                "recent": s.InspectionSerializer(
                    inspections.select_related("model", "job")[:6], many=True
                ).data,
            }
        )


class ReportViewSet(viewsets.ViewSet):
    def list(self, request):
        return Response(
            {
                "formats": ["xlsx", "pdf"],
                "resource": "inspections",
                "filters": ["from", "to"],
            }
        )

    @action(detail=False, methods=["get"])
    def export(self, request):
        from openpyxl import Workbook
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import (
            SimpleDocTemplate,
            Table,
            TableStyle,
            Paragraph,
            Spacer,
        )
        from reportlab.lib.styles import getSampleStyleSheet

        qs = owned(m.Inspection.objects.select_related("job", "model"), request.user)
        for param, field in [
            ("from", "created_at__date__gte"),
            ("to", "created_at__date__lte"),
        ]:
            if request.query_params.get(param):
                qs = qs.filter(
                    **{
                        field: serializers.DateField().run_validation(
                            request.query_params[param]
                        )
                    }
                )
        if qs.count() > 10000:
            raise ValidationError("Limita el rango a 10000 inspecciones.")
        rows = [
            [
                "ID",
                "Imagen",
                "Modelo",
                "Estado",
                "Lesiones",
                "Inferencia ms",
                "Total ms",
                "Fecha",
            ]
        ]
        for x in qs:
            rows.append(
                [
                    str(x.pk),
                    x.name,
                    x.model.name,
                    x.job.state,
                    len(x.detections) if x.job.state == "succeeded" else None,
                    x.inference_ms,
                    x.total_ms,
                    x.created_at.isoformat(),
                ]
            )
        out = io.BytesIO()
        fmt = request.query_params.get("format", "xlsx")
        if fmt == "xlsx":
            wb = Workbook()
            ws = wb.active
            ws.title = "Inspecciones"
            for row in rows:
                # Prevent spreadsheet formula injection from filenames or usernames.
                ws.append(
                    [
                        (
                            "'" + v
                            if isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
                            else v
                        )
                        for v in row
                    ]
                )
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 24
            wb.save(out)
        elif fmt == "pdf":
            styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(
                out, pagesize=landscape(A4), rightMargin=25, leftMargin=25
            )
            compact = [["Imagen", "Estado", "Lesiones", "Inferencia ms", "Fecha"]] + [
                [
                    str(r[1])[:35],
                    r[3],
                    r[4],
                    round(r[5], 2) if r[5] is not None else "—",
                    r[7][:10],
                ]
                for r in rows[1:]
            ]
            table = Table(compact, repeatRows=1, colWidths=[250, 100, 70, 110, 100])
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163d2e")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ]
                )
            )
            doc.build(
                [
                    Paragraph("HassLab · Reporte de inspecciones", styles["Title"]),
                    Paragraph(
                        "Sin detecciones no equivale a fruto sano. Los tiempos corresponden al servidor.",
                        styles["Normal"],
                    ),
                    Spacer(1, 16),
                    table,
                ]
            )
        else:
            raise ValidationError("Formato permitido: xlsx o pdf.")
        services.audit(
            request.user,
            "report.export",
            request.user,
            {"format": fmt, "rows": len(rows) - 1},
        )
        out.seek(0)
        return FileResponse(out, as_attachment=True, filename=f"inspecciones.{fmt}")


class AuditViewSet(ReadOnly):
    class AuditSerializer(serializers.ModelSerializer):
        class Meta:
            model = m.AuditEvent
            fields = "__all__"

    serializer_class = AuditSerializer

    def get_queryset(self):
        return owned(m.AuditEvent.objects.all(), self.request.user, "actor")


# Attach explicit contracts once, before router registration.
import sys
from .schema import annotate_views

annotate_views(sys.modules[__name__])
