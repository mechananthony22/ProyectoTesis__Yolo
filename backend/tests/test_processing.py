import io
import pytest
from PIL import Image
from unittest.mock import patch
from django.core.files.base import ContentFile
from apps.research import tasks,services
from apps.research.models import Job,ImageVariant,TrainingRun,AIModel,Inspection,Evaluation,EvaluationSample
from ai.preprocessing import transform
from ai.evaluation import classification,paired
from .helpers import png,dataset_version
pytestmark=pytest.mark.django_db

def trained_fixture(user):
    ds,version=dataset_version(user)
    job=Job.objects.create(owner=user,kind='training',state='succeeded')
    run=TrainingRun.objects.create(version=version,job=job)
    model=AIModel(run=run,name='TEST FIXTURE · no trained model',sha256='dummy')
    model.weights.save('fixture.pt',ContentFile(b'not-real-weights'))
    return ds,version,model

def test_letterbox_preserves_boxes_and_no_fruit_repetition():
    data=png(size=(100,50));box={'class_id':0,'x':.5,'y':.5,'w':1.,'h':1.}
    output,boxes=transform(data,[box],{'size':800})
    img=Image.open(io.BytesIO(output))
    assert img.size==(800,800) and img.getpixel((10,10))==(255,255,255)
    assert boxes[0]['x']==.5 and boxes[0]['y']==.5 and boxes[0]['w']==1 and boxes[0]['h']==.5

def test_crop_clips_boxes():
    boxes=[{'class_id':0,'x':.25,'y':.5,'w':.4,'h':.4},{'class_id':0,'x':.9,'y':.5,'w':.1,'h':.2}]
    output,result=transform(png(size=(100,100)),boxes,{'size':640,'crop':[0,0,.5,1]})
    assert len(result)==1 and result[0]['w']==pytest.approx(.4)

def test_preprocess_task_is_idempotent(user):
    ds,version=dataset_version(user);image=ds.images.first()
    job=Job.objects.create(owner=user,kind='preprocessing',config={'images':[str(image.pk)],'size':640,'crop':[0,0,1,1],'brightness':1,'contrast':1})
    before=image.file.read();image.file.close()
    tasks.preprocess(str(job.pk));job.refresh_from_db()
    assert job.state=='succeeded' and ImageVariant.objects.filter(job=job).count()==1
    tasks.preprocess(str(job.pk));assert ImageVariant.objects.filter(job=job).count()==1
    assert image.file.open("rb").read()==before
    image.file.close()

def test_failed_broker_retains_queued_job(user):
    job=Job.objects.create(owner=user,kind='preprocessing')
    with patch.object(tasks.preprocess,'apply_async',side_effect=ConnectionError):services.dispatch(job,tasks.preprocess)
    job.refresh_from_db();assert job.state=='queued' and 'Redis' in job.error

def test_training_api_rejects_unbounded_config(client,user):
    ds,version=dataset_version(user)
    r=client.post('/api/v1/training-runs/',{'version':str(version.pk),'architecture':'yolov8n.pt','epochs':10000},format='json')
    assert r.status_code==400

def test_training_api_creates_durable_job(client,user):
    ds,version=dataset_version(user)
    with patch.object(tasks.train,'apply_async',side_effect=ConnectionError):
        r=client.post('/api/v1/training-runs/',{'version':str(version.pk),'architecture':'yolov8n.pt','epochs':1},format='json')
    assert r.status_code==202 and r.data['job']['state']=='queued'

def test_model_activation_admin_only(client,user,admin):
    ds,version,model=trained_fixture(user)
    assert client.post(f'/api/v1/models/{model.pk}/activate/').status_code==403
    client.force_authenticate(admin)
    assert client.post(f'/api/v1/models/{model.pk}/activate/').status_code==200
    model.refresh_from_db();assert model.active

def test_inference_no_detections_is_success(user):
    import hashlib
    ds,version,model=trained_fixture(user)
    model.sha256=hashlib.sha256(b'not-real-weights').hexdigest();model.save()
    job=Job.objects.create(owner=user,kind='inference',config={'confidence':.25})
    inspection=Inspection(owner=user,model=model,job=job,name='fixture.png',sha256='fixture')
    inspection.original.save('fixture.png',ContentFile(png()))
    with patch('ai.engine.predict',return_value=([],12.0,png())):
        tasks.infer(str(job.pk))
    inspection.refresh_from_db();job.refresh_from_db()
    assert job.state=='succeeded' and inspection.detections==[] and inspection.review=='pending'

def test_report_formats_and_formula_injection(client,user):
    from openpyxl import load_workbook
    ds,version,model=trained_fixture(user)
    job=Job.objects.create(owner=user,kind='inference',state='failed')
    Inspection.objects.create(owner=user,model=model,job=job,name='=1+1.png',sha256='fixture')
    response=client.get('/api/v1/reports/export/?format=xlsx')
    assert response.status_code==200
    workbook=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
    assert workbook.active['B2'].value.startswith("'") and workbook.active['E2'].value is None
    response=client.get('/api/v1/reports/export/?format=pdf')
    assert response.status_code==200 and b''.join(response.streaming_content).startswith(b'%PDF')

def test_evaluation_blind_samples_export_and_duplicates(client,user):
    from openpyxl import load_workbook
    ds,version,model=trained_fixture(user)
    response=client.post('/api/v1/evaluations/',{'name':'Fixture','protocol':'Test only','version':str(version.pk),'model':str(model.pk),'confidence':.25},format='json')
    assert response.status_code==201
    eid=response.data['id']
    samples=client.get(f'/api/v1/evaluations/{eid}/samples/').data
    assert len(samples)==1 and 'truth' not in samples[0] and 'system_positive' not in samples[0]
    entry={'sample':str(samples[0]['id']),'participant':'P01','positive':True,'elapsed_ms':500}
    assert client.post(f'/api/v1/evaluations/{eid}/observe/',entry,format='json').status_code==201
    assert client.post(f'/api/v1/evaluations/{eid}/observe/',entry,format='json').status_code==400
    response=client.get(f'/api/v1/evaluations/{eid}/export/')
    assert response.status_code==200
    workbook=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
    assert workbook['Muestras'].max_row==2 and 'Protocolo' in workbook.sheetnames

def test_undefined_classification_metrics_are_null():
    result=classification([False,False],[False,False])
    assert result['precision'] is None and result['recall'] is None and result['accuracy']==1

def test_paired_analysis_uses_actual_observations():
    result=paired([True,False,True],[True,False,False],[False,False,True],[10,12,11],[500,600,550])
    assert result['system']['n']==3 and result['mcnemar_exact_p']==1.0 and result['mean_system_ms']==11
